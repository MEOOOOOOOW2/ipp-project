# report.py - adds Summary, Ratios, Scores, Charts and Checks sheets to the uploaded workbook
import io, os                               # in-memory files, environment settings
os.environ.setdefault("MPLCONFIGDIR", "/tmp")   # matplotlib needs a writable folder on Vercel
import matplotlib                           # plotting library
matplotlib.use("Agg")                       # draw to memory, no screen needed
import matplotlib.pyplot as plt             # plotting commands
import pandas as pd                         # missing-value checks
from openpyxl import load_workbook          # edits the original Excel file without losing its sheets
from openpyxl.drawing.image import Image    # lets us place chart pictures inside Excel
from engine.scoring import label, LIMITS    # final label and indicator names

def put(ws, df):                            # writes a table into a sheet
    ws.append(list(df.columns))             # header row
    for r in df.itertuples(index=False):    # every data row
        ws.append([None if pd.isna(v) else (v.item() if hasattr(v, "item") else v) for v in r])   # clean values

def picture(fig):                           # chart -> picture Excel can hold (small, to stay under the 4.5 MB limit)
    buf = io.BytesIO()                      # memory file
    fig.savefig(buf, format="png", dpi=70, bbox_inches="tight")   # save as PNG
    plt.close(fig)                          # free memory
    buf.seek(0)                             # rewind
    return Image(buf)                       # picture object

def build_workbook(data, t, info, missing, res):   # main writer: returns the finished Excel file as bytes
    wb = load_workbook(io.BytesIO(data))    # open the uploaded file, keeping original sheets
    last = res.iloc[-1]                     # latest window
    s = wb.create_sheet("Summary", 0)       # Summary goes first
    for k, v in [("Company", info["Company"]), ("Sector", info["Sector"]), ("Final label", label(res)),
                 ("Latest window", last["Window"]), ("Latest score", round(float(last["Score"]), 2)),
                 ("Latest band", last["Band"]), ("Red flags now", last["Which"])]:
        s.append([k, v])                    # one fact per row
    s.column_dimensions["A"].width = 18     # readable column widths
    s.column_dimensions["B"].width = 40
    ratios = wb.create_sheet("Ratios")      # actual ratio values for comparison across windows
    put(ratios, res[["Window"] + [c for c in res.columns if c.startswith("val: ")]].round(3))
    scores = wb.create_sheet("Scores")      # points and bands per window
    put(scores, res[["Window", "Points", "Score", "Band", "Which"] + [c for c in res.columns if c.startswith("pts: ")]].round(3))
    ch = wb.create_sheet("Charts")          # all charts on one sheet
    f1, a1 = plt.subplots(figsize=(7, 3))   # chart 1: score trend (line)
    a1.plot(res["Window"], res["Score"] * 100, marker="o")   # score per window
    for y in (20, 40, 60):                  # band boundaries
        a1.axhline(y, ls="--", color="grey", lw=0.8)
    a1.set_title("Composite red-flag score by window (%)"); plt.setp(a1.get_xticklabels(), rotation=45)
    ch.add_image(picture(f1), "A1")         # place chart 1
    f2, a2 = plt.subplots(figsize=(7, 3))   # chart 2: points in latest window (bar)
    names = list(LIMITS)                    # indicator names
    a2.bar(names, [last["pts: " + n] if pd.notna(last["pts: " + n]) else 0 for n in names])   # points 0/1/2
    a2.set_title(f"Points by indicator, {last['Window']}"); plt.setp(a2.get_xticklabels(), rotation=30, ha="right")
    ch.add_image(picture(f2), "A18")        # place chart 2
    f3, a3 = plt.subplots(figsize=(7, 3))   # chart 3: profit vs cash (line) - the core idea of the project
    a3.plot(t.index, t["Net profit"], marker="o", label="Net profit")     # reported profit
    a3.plot(t.index, t["CFO"], marker="o", label="Operating cash flow")   # cash actually earned
    a3.legend(); a3.set_title("Profit vs operating cash flow")
    ch.add_image(picture(f3), "A35")        # place chart 3
    chk = wb.create_sheet("Checks")         # data quality notes
    chk.append(["Years found", f"{t.index.min()} to {t.index.max()} ({len(t)} years)"])
    chk.append(["Fields not found", ", ".join(missing) or "none"])
    chk.append(["Note", "Indicators that need a missing field show N/A and are left out of the score"])
    out = io.BytesIO()                      # memory file for the result
    wb.save(out)                            # save workbook
    return out.getvalue()                   # bytes to send back
